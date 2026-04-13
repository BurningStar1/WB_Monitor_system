"""Insert section 3.3 text into VKR DOCX."""
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DOCX_SRC = r"D:\Dev\ВКР\ВКР_ПрепелицаПП_ИНБО-05.docx"
DOCX_OUT = r"D:\Dev\ВКР\ВКР_ПрепелицаПП_ИНБО-05_v2.docx"

doc = Document(DOCX_SRC)

# --- Find section 3.3 heading (idx 736) and section 4 heading ---
heading_elem = None
sec4_elem = None
for p in doc.paragraphs:
    if "3.3" in p.text and "емонстрац" in p.text:
        heading_elem = p._element
    if p.style.name == "Heading 1" and p.text.strip().startswith("4"):
        sec4_elem = p._element
        break

assert heading_elem is not None, "Section 3.3 not found"
assert sec4_elem is not None, "Section 4 not found"

# --- Remove empty paragraphs between 3.3 heading and section 4 ---
body = doc.element.body
sibling = heading_elem.getnext()
to_remove = []
while sibling is not None and sibling is not sec4_elem:
    if sibling.tag == qn("w:p"):
        text = sibling.text or ""
        # Also get text from runs
        full = "".join(t.text or "" for t in sibling.iter(qn("w:t")))
        if not full.strip():
            to_remove.append(sibling)
    sibling = sibling.getnext()

for elem in to_remove:
    body.remove(elem)

# --- Template: copy formatting from paragraph 700 (body text) ---
template_body = doc.paragraphs[700]._element
# Template for figure caption
fig_template = None
for p in doc.paragraphs:
    if p.style.name == "Для рисунков":
        fig_template = p._element
        break


def make_paragraph(text, style_name="ГОСТ_ОСН_ТЕКСТ"):
    """Create a new paragraph element with the given style and text."""
    tmpl = template_body if style_name == "ГОСТ_ОСН_ТЕКСТ" else fig_template
    new_p = deepcopy(tmpl)
    # Remove all existing runs
    for run in new_p.findall(qn("w:r")):
        new_p.remove(run)
    # Remove any bookmarkStart/End and other inline content
    for tag in ["w:bookmarkStart", "w:bookmarkEnd", "w:hyperlink"]:
        for el in new_p.findall(qn(tag)):
            new_p.remove(el)
    # Remove mc:AlternateContent (different namespace)
    MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    for el in new_p.findall(f"{{{MC_NS}}}AlternateContent"):
        new_p.remove(el)
    # Create new run
    run_el = OxmlElement("w:r")
    # Copy run properties from template
    if tmpl is not None:
        tmpl_runs = tmpl.findall(qn("w:r"))
        if tmpl_runs:
            rpr = tmpl_runs[0].find(qn("w:rPr"))
            if rpr is not None:
                run_el.append(deepcopy(rpr))
    t_el = OxmlElement("w:t")
    t_el.set(qn("xml:space"), "preserve")
    t_el.text = text
    run_el.append(t_el)
    new_p.append(run_el)
    # Set style
    ppr = new_p.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        new_p.insert(0, ppr)
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is None:
        pstyle = OxmlElement("w:pStyle")
        ppr.insert(0, pstyle)
    # Use internal style IDs, not display names
    STYLE_IDS = {
        "ГОСТ_ОСН_ТЕКСТ": "af7",
        "Для рисунков": "ae",
        "Normal": "Normal",
    }
    pstyle.set(qn("w:val"), STYLE_IDS.get(style_name, style_name))
    return new_p


# --- Content ---
BODY = "ГОСТ_ОСН_ТЕКСТ"
FIG = "Для рисунков"

content = [
    (BODY, (
        "Для подтверждения работоспособности разработанного аналитического "
        "сервиса выполнена демонстрация основных сценариев использования "
        "на реальных данных продавца маркетплейса Wildberries. Демонстрация "
        "проводилась на наборе данных за период с января по апрель 2026 г., "
        "включающем свыше 4 000 заказов, 3 000 продаж и 65 000 строк "
        "финансовых отчётов, загруженных из Statistics API и Report API "
        "маркетплейса."
    )),
    (BODY, (
        "Веб-интерфейс сервиса реализован средствами фреймворка Streamlit "
        "и включает главную страницу навигации и семь страниц интерактивных "
        "отчётов. Каждая страница снабжена боковой панелью фильтрации, "
        "позволяющей ограничить данные по периоду, бренду, предмету и "
        "артикулу поставщика. Результаты отображаются в виде KPI-карточек, "
        "интерактивных графиков Plotly и таблиц с возможностью сортировки."
    )),
    (BODY, (
        "На Рисунке 3.12 представлена главная страница приложения, "
        "содержащая перечень доступных отчётов с кратким описанием их "
        "назначения. Навигация между страницами осуществляется через "
        "боковое меню Streamlit."
    )),
    (FIG, "Рисунок 3.12 \u2014 Главная страница аналитического сервиса"),
    (BODY, (
        "Центральным элементом сервиса является KPI-дашборд, "
        "представленный на Рисунке 3.13. В верхней части страницы "
        "расположены четыре агрегирующие карточки: \u00abРеализация\u00bb "
        "(валовая выручка с разбивкой на продажи и возвраты), "
        "\u00abУслуги WB\u00bb (комиссия, логистика, реклама, прочие "
        "удержания с процентной долей от реализации), "
        "\u00abНалоги и затраты\u00bb (себестоимость, налог УСН, "
        "дополнительные расходы), \u00abОперационная прибыль\u00bb "
        "(с расчётом маржинальности и рентабельности). Во второй строке "
        "размещены карточки дневной динамики с мини-графиками "
        "(спарклайнами) за выбранный период: заказы, продажи, логистика, "
        "реклама и суммарные услуги. Для расходных метрик (логистика, "
        "реклама, услуги) применена инвертированная цветовая индикация: "
        "рост отображается красным, снижение \u2014 зелёным."
    )),
    (FIG, "Рисунок 3.13 \u2014 KPI-дашборд: карточки показателей и спарклайны дневной динамики"),
    (BODY, (
        "Ниже карточек расположены аналитические графики, показанные "
        "на Рисунке 3.14: столбчатая диаграмма динамики заказов, продаж "
        "и прибыли по месяцам; совмещённый график выручки с кривой "
        "маржинальности; а также детализация по дням. В нижней части "
        "дашборда представлены горизонтальные диаграммы топ-10 товаров "
        "по операционной прибыли в разрезе брендов, предметов и артикулов "
        "поставщика."
    )),
    (FIG, "Рисунок 3.14 \u2014 KPI-дашборд: аналитические графики и топ-10 товаров"),
    (BODY, (
        "Для оценки недельных тенденций предназначен еженедельный отчёт, "
        "представленный на Рисунке 3.15. Страница отображает столбчатую "
        "диаграмму выручки и прибыли по ISO-неделям, а также график "
        "динамики продаж и возвратов. Ниже приведена таблица с детализацией "
        "по каждой неделе, включающая показатели выручки, прибыли, "
        "количества заказов, продаж и возвратов."
    )),
    (FIG, "Рисунок 3.15 \u2014 Еженедельный отчёт"),
    (BODY, (
        "Отчёт по артикулам, представленный на Рисунке 3.16, позволяет "
        "оценить эффективность каждого товара. Страница включает "
        "KPI-карточки общего количества артикулов, суммарной выручки и "
        "прибыли. Основной контент \u2014 интерактивная таблица со "
        "столбцами: идентификатор товара (nm_id), артикул поставщика, "
        "предмет, бренд, количество продаж и возвратов, выручка, "
        "себестоимость и прибыль. Предусмотрена фильтрация по категории "
        "товара. Также отображается горизонтальная диаграмма топ-15 "
        "товаров по выручке."
    )),
    (FIG, "Рисунок 3.16 \u2014 Отчёт по артикулам"),
    (BODY, (
        "Страница остатков на складах, представленная на Рисунке 3.17, "
        "отображает текущий срез складских запасов. KPI-карточки показывают "
        "количество позиций, общий остаток и количество товаров в пути. "
        "Круговая диаграмма демонстрирует распределение остатков по складам "
        "Wildberries. Ниже приведена таблица с детализацией по каждому "
        "артикулу и складу, включающая количество на складе, в пути к "
        "клиенту и от клиента."
    )),
    (FIG, "Рисунок 3.17 \u2014 Отчёт по остаткам на складах"),
    (BODY, (
        "На Рисунке 3.18 представлена страница ABC-анализа, реализующая "
        "классификацию товаров по вкладу в выручку. В верхней части "
        "отображена сводка по категориям: количество товаров и доля выручки "
        "для групп A (80\u00a0% выручки), B (15\u00a0%) и C (5\u00a0%). "
        "Визуализация включает круговую диаграмму долей выручки по "
        "категориям и кривую Парето \u2014 график накопленной доли выручки "
        "с пороговыми линиями на уровнях 80\u00a0% и 95\u00a0%. Таблица "
        "детализации содержит каждый артикул с указанием категории ABC, "
        "выручки, прибыли и накопленной доли."
    )),
    (FIG, "Рисунок 3.18 \u2014 ABC-анализ товаров"),
    (BODY, (
        "Отчёт о прибыли, представленный на Рисунке 3.19, обеспечивает "
        "детализированный финансовый анализ. KPI-карточки отображают "
        "выручку, себестоимость, прибыль и маржинальность. Столбчатая "
        "диаграмма структуры финансового результата наглядно показывает "
        "соотношение компонентов: выручки, себестоимости, комиссий и "
        "прибыли. Линейный график демонстрирует тренд дневной прибыли за "
        "выбранный период. Детализированная таблица содержит данные по "
        "каждому артикулу и дню с расчётом маржи и ROI."
    )),
    (FIG, "Рисунок 3.19 \u2014 Отчёт о прибыли"),
    (BODY, (
        "Регламентный отчёт за период, представленный на Рисунке 3.20, "
        "формирует помесячную сводку ключевых показателей. В верхней части "
        "расположены KPI-карточки за последний доступный месяц. Столбчатая "
        "диаграмма отображает динамику выручки и прибыли по месяцам. "
        "Таблица детализации содержит помесячную разбивку всех финансовых "
        "показателей: выручка, комиссия, себестоимость, дополнительные "
        "расходы, налоги, прибыль и операционная прибыль."
    )),
    (FIG, "Рисунок 3.20 \u2014 Регламентный отчёт за период"),
    (BODY, (
        "Для проверки корректности расчётов была проведена сверка данных "
        "сервиса с официальным отчётом о прибылях и убытках (ОПИУ), "
        "формируемым маркетплейсом Wildberries. Сопоставление проводилось "
        "за январь 2026 г. по ключевым показателям. Результаты сверки "
        "приведены в Таблице 3.5."
    )),
]

# --- Insert paragraphs before section 4 ---
for style, text in content:
    new_elem = make_paragraph(text, style)
    sec4_elem.addprevious(new_elem)

# --- Table caption ---
caption = make_paragraph(
    "Таблица 3.5 \u2014 Сравнение показателей сервиса с ОПИУ Wildberries "
    "за январь 2026 г.",
    "Normal",
)
sec4_elem.addprevious(caption)

# --- Comparison table ---
tbl = OxmlElement("w:tbl")

tblPr = OxmlElement("w:tblPr")
tblStyle = OxmlElement("w:tblStyle")
tblStyle.set(qn("w:val"), "a5")
tblPr.append(tblStyle)
tblW = OxmlElement("w:tblW")
tblW.set(qn("w:w"), "0")
tblW.set(qn("w:type"), "auto")
tblPr.append(tblW)

tblBorders = OxmlElement("w:tblBorders")
for bname in ("top", "left", "bottom", "right", "insideH", "insideV"):
    b = OxmlElement(f"w:{bname}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "4")
    b.set(qn("w:space"), "0")
    b.set(qn("w:color"), "000000")
    tblBorders.append(b)
tblPr.append(tblBorders)
tbl.append(tblPr)

tblGrid = OxmlElement("w:tblGrid")
for _ in range(4):
    col = OxmlElement("w:gridCol")
    col.set(qn("w:w"), "2400")
    tblGrid.append(col)
tbl.append(tblGrid)

rows_data = [
    ["Показатель", "Сервис", "ОПИУ WB", "Отклонение"],
    ["Реализация до СПП", "2\u00a0616\u00a0134,23", "2\u00a0616\u00a0134,23", "0,00"],
    ["К перечислению", "1\u00a0723\u00a0337,70", "1\u00a0723\u00a0337,70", "0,00"],
    ["Логистика", "380\u00a0411,71", "380\u00a0411,71", "0,00"],
    ["Хранение", "4\u00a0019,79", "4\u00a0019,78", "0,01"],
]

for row_data in rows_data:
    tr = OxmlElement("w:tr")
    for cell_text in row_data:
        tc = OxmlElement("w:tc")
        p = OxmlElement("w:p")
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "24")
        rPr.append(sz)
        szCs = OxmlElement("w:szCs")
        szCs.set(qn("w:val"), "24")
        rPr.append(szCs)
        rFonts = OxmlElement("w:rFonts")
        rFonts.set(qn("w:ascii"), "Times New Roman")
        rFonts.set(qn("w:hAnsi"), "Times New Roman")
        rPr.append(rFonts)
        r.append(rPr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = cell_text
        r.append(t)
        p.append(r)
        tc.append(p)
        tr.append(tc)
    tbl.append(tr)

sec4_elem.addprevious(tbl)

# --- Concluding paragraphs ---
conclusions = [
    (
        "Как видно из Таблицы 3.5, ключевые финансовые показатели сервиса "
        "совпадают с данными официального отчёта Wildberries с точностью "
        "до копеек. Незначительное отклонение в показателе хранения "
        "(0,01 руб.) объясняется различием в порядке округления "
        "при агрегации."
    ),
    (
        "Таким образом, продемонстрировано, что разработанный аналитический "
        "сервис корректно загружает данные из API маркетплейса, формирует "
        "витрины данных и визуализирует интерактивные отчёты. Система "
        "обеспечивает прозрачность финансовых показателей: расчёты прибыли "
        "учитывают все категории расходов WB \u2014 комиссию, логистику, "
        "хранение, штрафы, приёмку, рекламу и прочие удержания, что "
        "позволяет получить результат, сопоставимый с официальной "
        "отчётностью маркетплейса. Полная видеодемонстрация работы "
        "сервиса приведена в Приложении Ж."
    ),
]

for text in conclusions:
    elem = make_paragraph(text, BODY)
    sec4_elem.addprevious(elem)

# --- Save ---
doc.save(DOCX_OUT)
print("Section 3.3 inserted successfully!")
print("  - 9 figure placeholders (3.12 — 3.20)")
print("  - 1 comparison table (Table 3.5)")
print("  - Reference to Appendix Ж (video demo)")
